import tempfile
import unittest
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from excel_report import build_report
from import_articles import read_articles_from_xls


class ImportArticlesTests(unittest.TestCase):
    def test_new_layout_uses_only_white_rows(self) -> None:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append([
            "Менеджер",
            "Бренд",
            "Артикул",
            "Название конкурента",
            "Мусор",
            "Kerry",
            "Название Kerry",
        ])
        sheet.append([
            "Рыбаков",
            "Abro",
            "SG330R",
            "Клей секундный ABRO",
            "",
            "КR-153-2",
            "Универсальный секундный клей",
        ])
        sheet.append([
            "Рыбаков",
            "Navr",
            "NV5055",
            "Супер клей 505",
            "",
            "KR-153-2",
            "Универсальный секундный клей",
        ])
        yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
        for cell in sheet[3]:
            cell.fill = yellow

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "articles.xlsx"
            workbook.save(path)
            articles = read_articles_from_xls(str(path))

        self.assertEqual(list(articles), ["KR-153-2"])
        self.assertEqual(articles["KR-153-2"]["brand"], "Kerry")
        self.assertEqual(
            articles["KR-153-2"]["competitors"],
            {
                "SG330R": {
                    "brand": "Abro",
                    "name": "Клей секундный ABRO",
                }
            },
        )


class ExcelReportTests(unittest.TestCase):
    def test_report_contains_competitor_product_name(self) -> None:
        report = build_report(
            rows=[{
                "brand": "Kerry",
                "article": "KR-153-2",
                "name": "Универсальный секундный клей",
                "our_price": 250,
                "competitors": {
                    "Abro": {
                        "price": 210,
                        "article": "SG330R",
                        "catalog": "",
                        "article_id": "",
                        "input_name": "Клей секундный ABRO",
                    }
                },
            }],
            all_brands=["Abro"],
        )

        workbook = openpyxl.load_workbook(BytesIO(report), data_only=False)
        self.assertEqual(
            workbook["Цены"].cell(row=2, column=7).value,
            "210₽ (SG330R) — Клей секундный ABRO",
        )
        self.assertEqual(
            workbook["Отсортированные"].cell(row=2, column=7).value,
            "210₽ (SG330R) — Abro — Клей секундный ABRO",
        )


if __name__ == "__main__":
    unittest.main()
